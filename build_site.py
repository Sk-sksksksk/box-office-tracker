#!/usr/bin/env python3
"""
Build the box office website from your tracker workbook.

    python build_site.py
    python build_site.py --file box_office_tracker.xlsx --out site

Reads the Films tab and the Daily Log tab of the workbook and writes:
    site/index.html   the finished website, with your data inside it (open it in a browser)
    site/data.json    the same data as plain JSON

Needs site_template.html in the same folder as this script.
Close the workbook in Excel first. Run this again whenever you add new figures.

What goes into the site
-----------------------
- Only rows with Gross Basis = "Net India" (the site compares like with like).
- "Trade estimate" rows become trade figures, one per source per day.
- "Producer-reported" rows become the producer claim for that day.
- "Confirmed actual" rows are skipped for now and counted in the summary below.
- The example film ("... (delete me)") is skipped unless you pass --include-example.
"""
import argparse
import json
import sys
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

PLACEHOLDER = "const INJECTED=null;"


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def main():
    ap = argparse.ArgumentParser(description="Build the website from the tracker workbook.")
    ap.add_argument("--file", default="box_office_tracker.xlsx")
    ap.add_argument("--out", default="site", help="output folder (default: site)")
    ap.add_argument("--include-example", action="store_true", help="keep the sample film")
    args = ap.parse_args()

    here = Path(__file__).resolve().parent
    template = here / "site_template.html"
    if not template.exists():
        sys.exit(f"Error: {template.name} must be in the same folder as this script.")
    if not Path(args.file).exists():
        sys.exit(f"Error: workbook not found: {args.file}")

    try:
        wb = load_workbook(args.file, data_only=False)
    except PermissionError:
        sys.exit("Error: could not open the workbook. Close it in Excel and try again.")

    # Films tab
    films = {}
    for r in range(2, 101):
        name, lang, rel = (wb["Films"].cell(row=r, column=c).value for c in (1, 2, 3))
        if not name:
            continue
        if "(delete me)" in str(name) and not args.include_example:
            continue
        rd = as_date(rel)
        if rd is None:
            print(f"Skipped film '{name}': release date is missing or unreadable.")
            continue
        films[name] = {"name": str(name), "lang": str(lang or "Other"), "rel": rd.isoformat(), "days": {}}

    # Daily Log tab (input columns only; formula columns are ignored)
    log = wb["Daily Log"]
    used = {"trade": 0, "producer": 0}
    skipped = {"basis": 0, "confirmed": 0, "unknown film": 0, "bad row": 0}
    for r in range(2, 501):
        film, day, gross, basis, source, ftype = (
            log.cell(row=r, column=c).value for c in (2, 3, 5, 6, 7, 8))
        if not film and gross in (None, ""):
            continue
        if film not in films:
            if film and "(delete me)" not in str(film):
                skipped["unknown film"] += 1
            continue
        d = as_date(day)
        if d is None or not isinstance(gross, (int, float)) or not source or not ftype:
            skipped["bad row"] += 1
            continue
        if basis != "Net India":
            skipped["basis"] += 1
            continue
        if ftype == "Confirmed actual":
            skipped["confirmed"] += 1
            continue
        entry = films[film]["days"].setdefault(d.isoformat(), {"tr": {}, "ps": []})
        if ftype == "Trade estimate":
            entry["tr"][str(source)] = float(gross)
            used["trade"] += 1
        elif ftype == "Producer-reported":
            entry["ps"].append(float(gross))
            used["producer"] += 1

    out_films = []
    for f in films.values():
        days = []
        for iso in sorted(f["days"]):
            e = f["days"][iso]
            days.append({"date": iso, "tr": e["tr"],
                         "p": round(sum(e["ps"]) / len(e["ps"]), 2) if e["ps"] else None})
        if days:
            out_films.append({"name": f["name"], "lang": f["lang"], "rel": f["rel"], "days": days})

    payload = {"generated": date.today().strftime("%d %b %Y"), "films": out_films}
    html = template.read_text(encoding="utf-8")
    if PLACEHOLDER not in html:
        sys.exit("Error: site_template.html does not contain the data placeholder.")
    blob = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    html = html.replace(PLACEHOLDER, f"const INJECTED={blob};", 1)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "index.html").write_text(html, encoding="utf-8")
    (out / "data.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Built {out / 'index.html'}: {len(out_films)} film(s), "
          f"{used['trade']} trade figure(s), {used['producer']} producer figure(s).")
    notes = [f"{n} {k}" for k, n in skipped.items() if n]
    if notes:
        print("Skipped rows: " + ", ".join(notes) + ".")
    if not out_films:
        print("No films with data yet. Add figures to the Daily Log first.")


if __name__ == "__main__":
    main()
