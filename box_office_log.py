#!/usr/bin/env python3
"""
Box office tracker: add rows to the Daily Log of box_office_tracker.xlsx.

Usage examples
--------------
Add one figure:
    python box_office_log.py add --film "Film Name" --gross 21.5 --source "Sacnilk"

Add a figure for an earlier day, from a producer poster:
    python box_office_log.py add --film "Film Name" --gross 26 --source "Producer / Official poster" \
        --type "Producer-reported" --date 2026-10-02

Import many rows from a CSV (columns: film,date,gross,source,basis,type,notes):
    python box_office_log.py import entries.csv

Check whether a site's robots.txt permits automated access to a URL:
    python box_office_log.py check-robots https://example.com/some/page

Close the workbook in Excel before running. The script only writes the blue input
columns and never touches the formula columns, so the flags and Summary keep working.
Formulas recalculate the next time you open the file in Excel.
"""
import argparse
import csv
import shutil
import sys
from datetime import datetime, date
from pathlib import Path
from urllib import robotparser
from urllib.parse import urlparse

from openpyxl import load_workbook

DEFAULT_FILE = "box_office_tracker.xlsx"
LOG = "Daily Log"
FIRST_ROW, LAST_ROW = 2, 500          # must match the workbook layout
# Daily Log input columns (A..I). Formula columns D and J-N are never written.
C_LOGGED, C_FILM, C_DATE, C_GROSS, C_BASIS, C_SOURCE, C_TYPE, C_NOTES = 1, 2, 3, 5, 6, 7, 8, 9


class TrackerError(Exception):
    pass


def parse_date(text):
    if isinstance(text, (datetime, date)):
        return datetime(text.year, text.month, text.day)
    text = str(text).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise TrackerError(f"Could not read date '{text}'. Use YYYY-MM-DD.")


def column_values(ws, col, first=2, last=100):
    return [ws.cell(row=r, column=col).value for r in range(first, last + 1)
            if ws.cell(row=r, column=col).value not in (None, "")]


class Tracker:
    def __init__(self, path):
        self.path = Path(path)
        if not self.path.exists():
            raise TrackerError(f"File not found: {self.path}")
        self.wb = load_workbook(self.path)  # keeps formulas, dropdowns, formatting
        self.log = self.wb[LOG]
        films = self.wb["Films"]
        self.release = {films.cell(row=r, column=1).value: films.cell(row=r, column=3).value
                        for r in range(2, 101) if films.cell(row=r, column=1).value}
        self.sources = column_values(self.wb["Sources"], 1, 2, 50)
        lists = self.wb["Lists"]
        self.bases = column_values(lists, 1, 2, 20)
        self.types = column_values(lists, 2, 2, 20)

    def _find_existing(self, film, day, source, basis):
        for r in range(FIRST_ROW, LAST_ROW + 1):
            row_date = self.log.cell(row=r, column=C_DATE).value
            if (self.log.cell(row=r, column=C_FILM).value == film
                    and row_date is not None and parse_date(row_date) == day
                    and self.log.cell(row=r, column=C_SOURCE).value == source
                    and self.log.cell(row=r, column=C_BASIS).value == basis):
                return r
        return None

    def _first_empty_row(self):
        for r in range(FIRST_ROW, LAST_ROW + 1):
            if (self.log.cell(row=r, column=C_FILM).value in (None, "")
                    and self.log.cell(row=r, column=C_GROSS).value in (None, "")):
                return r
        raise TrackerError(f"Daily Log is full ({LAST_ROW} rows). Copy the file and start a new one.")

    def validate(self, film, day, gross, source, basis, ftype):
        if film not in self.release:
            raise TrackerError(f"Film '{film}' is not on the Films tab. Add it there first. "
                               f"Known films: {', '.join(self.release) or 'none'}")
        if source not in self.sources:
            raise TrackerError(f"Source '{source}' is not on the Sources tab. Add it there first. "
                               f"Known sources: {', '.join(self.sources)}")
        if basis not in self.bases:
            raise TrackerError(f"Basis must be one of: {', '.join(self.bases)}")
        if ftype not in self.types:
            raise TrackerError(f"Type must be one of: {', '.join(self.types)}")
        try:
            gross = float(gross)
        except (TypeError, ValueError):
            raise TrackerError(f"Gross '{gross}' is not a number.")
        if gross <= 0:
            raise TrackerError("Gross must be greater than 0 (Rs crore).")
        rel = self.release[film]
        if rel is not None and day < parse_date(rel):
            raise TrackerError(f"Date {day:%Y-%m-%d} is before the release date {parse_date(rel):%Y-%m-%d}.")
        return gross

    def add(self, film, day, gross, source, basis="Net India", ftype="Trade estimate", notes=""):
        gross = self.validate(film, day, gross, source, basis, ftype)
        existing = self._find_existing(film, day, source, basis)
        now = datetime.now().replace(microsecond=0)
        if existing:
            # Update in place so the formulas do not double-count this source.
            old = self.log.cell(row=existing, column=C_GROSS).value
            self.log.cell(row=existing, column=C_GROSS, value=gross)
            self.log.cell(row=existing, column=C_LOGGED, value=now)
            self.log.cell(row=existing, column=C_TYPE, value=ftype)
            note = f"revised from {old}" + (f"; {notes}" if notes else "")
            self.log.cell(row=existing, column=C_NOTES, value=note)
            return f"Updated row {existing}: {film}, {day:%d-%b-%Y}, {source}: {old} -> {gross}"
        r = self._first_empty_row()
        values = {C_LOGGED: now, C_FILM: film, C_DATE: day, C_GROSS: gross,
                  C_BASIS: basis, C_SOURCE: source, C_TYPE: ftype, C_NOTES: notes or None}
        for col, v in values.items():
            self.log.cell(row=r, column=col, value=v)
        return f"Added row {r}: {film}, {day:%d-%b-%Y}, {source}, Rs {gross} cr ({basis}, {ftype})"

    def save(self):
        backup = self.path.with_suffix(".backup.xlsx")
        shutil.copy2(self.path, backup)       # one-step undo if something goes wrong
        try:
            self.wb.save(self.path)
        except PermissionError:
            raise TrackerError("Could not save. Close the workbook in Excel and try again.")


def robots_check(url):
    parts = urlparse(url)
    if not parts.scheme or not parts.netloc:
        raise TrackerError("Give a full URL, e.g. https://example.com/page")
    rp = robotparser.RobotFileParser()
    rp.set_url(f"{parts.scheme}://{parts.netloc}/robots.txt")
    try:
        rp.read()
    except Exception as e:
        raise TrackerError(f"Could not read robots.txt ({e}). Treat access as NOT confirmed.")
    allowed = rp.can_fetch("*", url)
    print(f"robots.txt {'allows' if allowed else 'DISALLOWS'} automated access to {url}")
    print("Note: robots.txt is only one signal. Also read the site's terms of service, "
          "and record your finding on the Sources tab.")


def main():
    p = argparse.ArgumentParser(description="Add box office figures to the tracker workbook.")
    p.add_argument("--file", default=DEFAULT_FILE, help=f"workbook path (default {DEFAULT_FILE})")
    sub = p.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("add", help="add one figure")
    a.add_argument("--film", required=True)
    a.add_argument("--gross", required=True, type=float, help="Rs crore")
    a.add_argument("--source", required=True)
    a.add_argument("--date", help="business date YYYY-MM-DD (default: today)")
    a.add_argument("--basis", default="Net India")
    a.add_argument("--type", default="Trade estimate", dest="ftype")
    a.add_argument("--notes", default="")

    i = sub.add_parser("import", help="import rows from a CSV")
    i.add_argument("csv_file")

    r = sub.add_parser("check-robots", help="check robots.txt for a URL")
    r.add_argument("url")

    args = p.parse_args()
    try:
        if args.cmd == "check-robots":
            robots_check(args.url)
            return
        t = Tracker(args.file)
        if args.cmd == "add":
            day = parse_date(args.date) if args.date else parse_date(date.today())
            print(t.add(args.film, day, args.gross, args.source, args.basis, args.ftype, args.notes))
            t.save()
        else:
            ok = bad = 0
            with open(args.csv_file, newline="", encoding="utf-8-sig") as f:
                for n, row in enumerate(csv.DictReader(f), start=2):
                    try:
                        msg = t.add(
                            row["film"].strip(), parse_date(row["date"]), row["gross"],
                            row["source"].strip(),
                            (row.get("basis") or "Net India").strip(),
                            (row.get("type") or "Trade estimate").strip(),
                            (row.get("notes") or "").strip())
                        print(msg); ok += 1
                    except (TrackerError, KeyError) as e:
                        print(f"CSV line {n} skipped: {e}"); bad += 1
            if ok:
                t.save()
            print(f"Done: {ok} saved, {bad} skipped.")
    except TrackerError as e:
        sys.exit(f"Error: {e}")


if __name__ == "__main__":
    main()
